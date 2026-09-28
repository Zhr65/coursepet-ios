# MARK: - 课件文件解析与分块（模式 7 RAG 的文件级入口）
# 职责：iOS 上传的 PDF/DOCX/PPTX/TXT → 抽取纯文本 → 语义分块（滑窗 + 句界优先）。
# 解析库全为纯 Python（pypdf / python-docx / python-pptx），VM 上 pip 即装，无系统依赖。
# 设计约束：
#   - PG 不接受 \x00 字节，PDF/DOCX 抽出的文本必须清洗（清洗函数里统一处理）
#   - 扫描版 PDF 抽不出文字：诚实抛错，绝不静默入库空块
#   - 中文长度用 len()（码点数），不按字节
import io
import re

MAX_CHARS = 60_000   # 单文件抽取文本上限（超出截断，防巨型课件拖垮嵌入与存储）
MAX_CHUNKS = 40      # 单文件分块上限（一块一行 course_docs，防刷表）
CHUNK_SIZE = 500     # 目标块长（字符）
CHUNK_OVERLAP = 80   # 相邻块重叠（保住跨块的上下文连续性）

_SENTENCE_END = "。！？\n；"  # 优先切句边界；找不到才硬切

_ZERO_WIDTH_RE = re.compile(r"[\u200b-\u200f\u2060\ufeff]")
_MULTI_BLANK_RE = re.compile(r"\n{3,}")


def clean_text(raw: str) -> str:
    """抽取文本的统一清洗：去 NUL/零宽字符、统一换行、压掉 3+ 连续空行。"""
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x00", "")
    text = _ZERO_WIDTH_RE.sub("", text)
    text = _MULTI_BLANK_RE.sub("\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> str:
    """按扩展名分派解析；返回清洗后的纯文本（空文本/不支持的格式抛 ValueError）。"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if ext == "pdf":
            text = _extract_pdf(data)
        elif ext == "txt" or ext == "md":
            text = _extract_txt(data)
        elif ext == "docx":
            text = _extract_docx(data)
        elif ext == "pptx":
            text = _extract_pptx(data)
        else:
            raise ValueError(f"暂不支持的文件格式：.{ext}（支持 pdf / docx / pptx / txt / md）")
    except ValueError:
        raise
    except Exception:
        raise ValueError("文件解析失败：文件可能已损坏或加密")
    text = clean_text(text)[:MAX_CHARS]
    if len(text) < 20:
        raise ValueError("没有从文件里提取到足够的文字（可能是扫描/图片版 PDF，需要 OCR 才能识别）")
    return text


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(p for p in pages if p)


def _extract_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("gbk", errors="ignore")


def _extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs)
                    if t.strip():
                        lines.append(t)
    return "\n".join(lines)


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP,
               max_chunks: int | None = MAX_CHUNKS) -> list[str]:
    """语义滑窗分块：优先在句末标点/换行处切，切点不足时硬切；块间保留 overlap 重叠。

    max_chunks=None 时不设上限（truncated 判定用），默认按 MAX_CHUNKS 截断。"""
    if len(text) <= size:
        chunks = [text]
    else:
        chunks = []
        start = 0
        while start < len(text):
            window = text[start:start + size]
            if start + size >= len(text):       # 剩余不足一块：整段收尾
                chunks.append(text[start:])
                break
            # 在窗口后半段找最后一个句界（前 250 字硬切边界内不找，避免块过短）
            cut = -1
            for i in range(len(window) - 1, size // 2 - 1, -1):
                if window[i] in _SENTENCE_END:
                    cut = i + 1
                    break
            if cut <= 0:
                cut = size
            chunks.append(text[start:start + cut].strip())
            start = start + cut - overlap
            if start < 0:
                start = 0
        chunks = [c for c in chunks if c]
    # 末块太短（<30 字）并入前块，避免检索召回一段废话
    if len(chunks) >= 2 and len(chunks[-1]) < 30:
        chunks[-2] += chunks[-1]
        chunks.pop()
    if max_chunks is not None and len(chunks) > max_chunks:
        chunks = chunks[:max_chunks]
    return chunks
